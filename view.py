import csv
import tkinter
from tkinter import *
from tkinter import ttk
from tkinter import filedialog
import os
import pygame as pygame
from pubsub import pub
from PIL import Image, ImageTk
import time
from mutagen.mp3 import MP3
import tkinter.ttk as ttk
from tkinter import colorchooser
import tkinter.messagebox as mb
from openpyxl.workbook import Workbook
from openpyxl import load_workbook
import pandas as pd
from google.cloud import speech_v1p1beta1 as speech

pygame.mixer.init()


class View:
    def __init__(self, container, controller):
        self.container = container
        self.controller = controller
        self.annotation_box = None
        # super().__init__(container)
        # self.controller = controller
        self.setup()


    def setup(self):
        #this class will run first
        #calling the methos to setup the user interface
        self.create_menu()
        self.create_panedwindow()
        self.create_audio_frame()
        self.create_transcript_frame()
        self.create_annotation_frame()
        # self.create_file_option()
        self.setup_menu()
        self.setup_panedwindow()
        self.setup_audio_frame()
        self.setup_transcript_frame()
        self.setup_annotation_frame()

    def create_menu(self):
        self.frame = Frame(self.container)
        self.menuBar = Menu(self.frame)
        self.home = Menu(self.menuBar, tearoff=0, activebackground='green', activeforeground='white')
        self.menuBar.add_cascade(label='IE Tool', menu=self.home)
        self.home.add_command(label='Industry')
        self.home.add_command(label='Research')
        self.home.add_separator()
        self.home.add_command(label='Exit', command=self.container.quit)

    def create_panedwindow(self):
        self.pw = PanedWindow(self.container, orient=HORIZONTAL)
        self.leftFrame = Frame(self.pw)
        self.rightFrame = Frame(self.pw)
        self.annotationFrame = Frame(self.pw)

    def create_audio_frame(self):
        self.upload_button = Button(self.leftFrame, text='Upload Audio File', command=lambda: self.uploadAction())
        self.remove_button = Button(self.leftFrame, text='Remove Audio File', command=lambda: self.removeAudio())
        self.audio_box = Listbox(self.leftFrame, bg='white', fg='black', width=60)

    def create_transcript_frame(self):
        script_dir = os.path.dirname(__file__)  # <-- absolute dir the script is in
        rel_path_play = "images/play_pause.png"
        rel_path_stop = "images/stop.png"
        rel_path_forward = "images/forward.png"
        rel_path_rewind = "images/rewind.png"
        abs_file_path_play = os.path.join(script_dir, rel_path_play)
        abs_file_path_stop = os.path.join(script_dir, rel_path_stop)
        abs_file_path_forward = os.path.join(script_dir, rel_path_forward)
        abs_file_rewind = os.path.join(script_dir, rel_path_rewind)
        self.play_pause_btn_img = Image.open(abs_file_path_play)
        self.stop_btn_img = Image.open(abs_file_path_stop)
        self.forward_btn_img = Image.open(abs_file_path_forward)
        self.rewind_btn_img = Image.open(abs_file_rewind)

        # Resize the Image
        self.resized_play_pause_img = self.play_pause_btn_img.resize((50, 50), Image.ANTIALIAS)
        self.resized_play_pausebtn_img = ImageTk.PhotoImage(self.resized_play_pause_img)

        self.resized_stop_img = self.stop_btn_img.resize((50, 50), Image.ANTIALIAS)
        self.resized_stopbtn_img = ImageTk.PhotoImage(self.resized_stop_img)

        self.resized_forward_img = self.forward_btn_img.resize((50, 50), Image.ANTIALIAS)
        self.resized_forwardbtn_img = ImageTk.PhotoImage(self.resized_forward_img)

        self.resized_rewind_img = self.rewind_btn_img.resize((50, 50), Image.ANTIALIAS)
        self.resized_rewindbtn_img = ImageTk.PhotoImage(self.resized_rewind_img)

        # label
        self.audio_Label = Label(self.rightFrame, text='')

        # Create Status Bar
        self.status_bar = Label(self.rightFrame, text='', bd=1, relief=GROOVE, anchor=E)

        # Create Player Control Frame
        self.control_frame = Frame(self.rightFrame)

        self.play_pause_btn = Button(self.control_frame, image=self.resized_play_pausebtn_img, borderwidth=0, command=lambda: self.play_pause())
        self.stop_btn = Button(self.control_frame, image=self.resized_stopbtn_img, borderwidth=0, command=lambda: self.stop_audio())
        self.forward_btn = Button(self.control_frame, image=self.resized_forwardbtn_img, borderwidth=0)
        self.rewind_btn = Button(self.control_frame, image=self.resized_rewindbtn_img, borderwidth=0)

        self.slider = ttk.Scale(self.rightFrame, from_=0, to=100, orient=HORIZONTAL, value=0, length=360, command=self.slide)

        self.transcribe_button = Button(self.rightFrame, text='Transcribe', command=lambda: (self.transcribe_action(),self.get_str_val()))

        self.transcribe_file = Button(self.rightFrame, text='File', command=lambda: self.create_file_option())

        self.text_frame = Frame(self.rightFrame)

        self.text_scroll = Scrollbar(self.text_frame)

        self.text_box = Text(self.text_frame, width=40, height=15, font=('times new roman', 16), selectbackground='yellow',
                        selectforeground='black', yscrollcommand=self.text_scroll.set)
        self.str_val = tkinter.StringVar()
        # # self.r = Button(self.text_box, textvariable=self.str_val, background='Green',
        #            command=lambda m=self.str_val: self.controller.timelink(m))

    def create_file_option(self):
        self.top = Toplevel(self.container)
        self.top.title("File Menu")
        self.top.geometry("150x150")
        self.open_btn = Button(self.top, text="Open", command=lambda: self.open_file())
        self.open_btn.pack(pady=5)
        self.save_btn = Button(self.top, text="Save", command=lambda: self.save_file())
        self.save_btn.pack(pady=5)
        self.saveAs_btn = Button(self.top, text="SaveAs", command=lambda: self.saveAs_file())
        self.saveAs_btn.pack(pady=5)

    def create_annotation_frame(self):
        self.tool_bar = Frame(self.annotationFrame)
        self.highlight_button = Button(self.tool_bar, text='Highlight', command=lambda: self.highlight())
        self.clear_highlight_button = Button(self.tool_bar, text='Clear Highlight', command=lambda: self.clear_highlight())
        self.extract_button = Button(self.tool_bar, text='Extract', command=lambda: self.gettext())
        self.delete_annotation_button = Button(self.tool_bar, text='Remove Annotation', command=lambda: self.delete_annotation())
        self.annotation_file = Button(self.tool_bar, text='File', command=lambda: self.create_annotation_file_option())
        self.annotation_search = Button(self.tool_bar, text='Search', command=lambda: self.create_searchRecords())
        self.s = ttk.Style()
        self.annotation_box = ttk.Treeview(self.annotationFrame, column=("c1", "c2", "c3"), show='headings', height=5)
        self.annotation_box.column("# 1", anchor=CENTER)
        self.annotation_box.heading("# 1", text="Sentence")
        self.annotation_box.column("# 2", anchor=CENTER)
        self.annotation_box.heading("# 2", text="Annotation Tag")
        self.annotation_box.column("# 3", anchor=CENTER)
        self.annotation_box.heading("# 3", text="Description")
        # annotation_box = Listbox(annotationFrame, bg='white', fg='black', width=60)
        self.addannotate_Tag = Button(self.annotationFrame, text='Add Tag', command=lambda: self.add_Tag())

    def create_annotation_file_option(self):
        self.listtop = Toplevel(self.container)
        self.listtop.title("File Menu")
        self.listtop.geometry("150x150")
        self.open_list_btn = Button(self.listtop, text="Open", command=lambda: self.open_excel_file())
        self.open_list_btn.pack(pady=5)
        self.save_list_btn = Button(self.listtop, text="Save", command=lambda: self.save_excel_file())
        self.save_list_btn.pack(pady=5)
        self.clear_list_btn = Button(self.listtop, text="Clear", command=lambda: self.clear_table())
        self.clear_list_btn.pack(pady=5)

    def setup_menu(self):
        self.container.config(menu=self.menuBar)
        self.frame.pack()

    def setup_panedwindow(self):
        self.pw.pack(fill=BOTH, expand=True)
        self.pw.configure(sashrelief=RAISED)
        self.pw.add(self.leftFrame, width=150)
        self.pw.add(self.rightFrame, width=500)
        self.pw.add(self.annotationFrame)

    def setup_audio_frame(self):
        self.upload_button.pack(pady=10)
        self.remove_button.pack(pady=10)
        self.audio_box.pack()

    def setup_transcript_frame(self):
        self.audio_Label.pack()
        self.status_bar.pack(fill=X, side=BOTTOM, ipady=1)
        self.control_frame.pack(pady=10)
        self.play_pause_btn.grid(row=0, column=0, padx=10)
        self.stop_btn.grid(row=0, column=2, padx=10)
        self.forward_btn.grid(row=0, column=3, padx=10)
        self.rewind_btn.grid(row=0, column=4, padx=10)
        self.slider.pack(pady=10)
        self.transcribe_button.pack(pady=10)
        self.transcribe_file.pack()
        self.text_frame.pack(pady=10)
        self.text_scroll.pack(side=RIGHT, fill=Y)
        self.text_box.pack(pady=10)

    def setup_annotation_frame(self):
        self.tool_bar.pack()
        self.highlight_button.grid(row=0, column=0, sticky=W, padx=5)
        self.clear_highlight_button.grid(row=0, column=1, sticky=W, padx=5)
        self.extract_button.grid(row=0, column=2, sticky=W, padx=5)
        self.delete_annotation_button.grid(row=0, column=3, sticky=W, padx=5)
        self.annotation_file.grid(row=0, column=4, sticky=W, padx=5)
        self.annotation_search.grid(row=0, column=5, sticky=W, padx=5)
        self.s.theme_use('clam')
        self.annotation_box.pack(pady=20)
        self.addannotate_Tag.pack(pady=20)

    def uploadAction(self):
        # global path
        audios = filedialog.askopenfilenames()
        self.controller.uploadAction(audios)

    def add_audio(self, audio):
        self.audio_box.insert(END, audio)

    def removeAudio(self):
        self.audio_box.delete(ANCHOR)
        pygame.mixer.music.stop()

    # create Global Pause Variable
    global paused
    paused = False

    global play
    play = False

    def play_pause(self):
        global play
        play = not play
        if play:
            self.controller.play_audio()
        else:
            self.controller.pause_audio(paused)

    def slide(self, x):
        # slider_label.config(text=f'{int(slider.get())} of {int(audio_total_length)}')
        self.audio = self.audio_box.get(ACTIVE)
        self.audio = f'{self.controller.uploadAction.path}{self.audio}'

        pygame.mixer.music.load(self.audio)
        pygame.mixer.music.play(loops=0, start=int(self.slider.get()))

    def stop_audio(self):
        self.status_bar.config(text='')
        self.slider.config(value=0)
        pygame.mixer.music.stop()
        self.audio_box.selection_clear(ACTIVE)

        self.status_bar.config(text='')
        self.audio_Label.config(text='')

        self.stopped = True

    def transcribe_action(self):
        self.controller.transcribe_action()

    def add_script(self, script):
        self.text_box.insert(END, script)

    def clear_textbox(self):
        self.text_box.delete("1.0", "end")

    def get_str_val(self):
        self.controller.str_val.get()

    def open_file(self):
        self.text_box.delete("1.0", END)
        self.transcribe_text_file = filedialog.askopenfilenames()
        self.controller.open_file()

    def add_text(self, transcribed_text):
        self.text_box.insert(END, transcribed_text)

    def save_file(self):
        self.controller.save_file()

    def saveAs_file(self):
        self.transcribe_file_save = filedialog.asksaveasfilename(defaultextension=".*",
                                                            filetypes=(("Text File", "*.txt"), ("Word File", "*.docx")))
        self.controller.saveAs_file()

    # global highlightCount
    # highlightCount = 0

    def highlight(self):
        self.highlightCount = 0
        my_color = colorchooser.askcolor()[1]
        self.highlightCount += 1
        try:
            self.text_box.tag_add(self.highlightCount, "sel.first", "sel.last")
            self.text_box.tag_configure(self.highlightCount, foreground="Black", background=my_color)
        except tkinter.TclError:
            pass

    def clear_highlight(self):
        self.text_box.tag_remove(self.highlightCount, "sel.first", "sel.last")

    global extracted_text

    def gettext(self):
        global extracted_text
        extracted_text = self.text_box.get("sel.first", "sel.last")
        print(extracted_text)
        self.annotation_box.insert('', 'end', values=(extracted_text, '', ''))

    def delete_annotation(self):
        selected_item = self.annotation_box.selection()[0]
        self.annotation_box.delete(selected_item)

    # def close_win(self, top):
    #     top.destroy()

    def open_excel_file(self):
        self.excel_file_name = filedialog.askopenfilenames()
        self.controller.open_list()

    def save_excel_file(self):
        self.save_annotation_list = filedialog.asksaveasfilename(defaultextension=".*",
                                                            filetypes=(
                                                                ("CSV File", "*.csv"), ("Word File", "*.docx")))
        self.controller.save_list()

    def clear_table(self):
        self.controller.clear_list()

    def create_searchRecords(self):
        # global search_entry, search
        self.search = Toplevel(self.container)
        self.search.title('Lookup Records')
        self.search.geometry('400x200')

        self.search_frame = LabelFrame(self.search, text='Annotation Name')
        self.search_entry = Entry(self.search_frame)
        self.search_button = Button(self.search_frame, text='Search Annotation', command=self.controller.lookup)
        self.search_frame.pack(padx=10, pady=10)
        self.search_entry.pack(padx=20, pady=20)
        self.search_button.pack(padx=20, pady=20)

    def add_Tag(self):
        self.tag_window = Toplevel(self.container)
        self.tag_window.geometry("750x400")
        self.tag_label = Label(self.tag_window, text='Select the Tag')
        self.tag_label.pack(pady=10)
        self.tags = ['Requirement', 'Follow up', 'Mistake', 'Custom']
        self.list_frame = Frame(self.tag_window)
        self.list_frame.pack()

        self.list_scroll = Scrollbar(self.list_frame)
        self.list_scroll.pack(side=RIGHT, fill=Y)

        self.dropdown_label = Text(self.list_frame, width=20, height=1)
        self.dropdown_label.pack()

        self.dropdown = Listbox(self.list_frame, selectmode="multiple", yscrollcommand=self.list_scroll.set)
        self.dropdown.bind('<<ListboxSelect>>', self.selectedValue)
        self.dropdown.pack(pady=5)

        for self.each_item in range(len(self.tags)):
            self.dropdown.insert(END, self.tags[self.each_item])
            self.dropdown.itemconfig(self.each_item)

        self.description_label = Label(self.tag_window, text='Description')
        self.description_label.pack(pady=5)

        self.description_frame = Frame(self.tag_window)
        self.description_frame.pack()

        self.description_scroll = Scrollbar(self.description_frame)
        self.description_scroll.pack(side=RIGHT, fill=Y)

        self.description_text = Text(self.description_frame, width=60, height=3, font=('times new roman', 16),
                                     selectbackground='yellow', selectforeground='black',
                                     yscrollcommand=self.text_scroll.set)
        self.description_text.pack(pady=10)

        self.button = Button(self.tag_window, text="Ok", command=lambda: [self.gettag_values(), self.tag_window.destroy()])
        self.button.pack(pady=10)

    def gettag_values(self):
        self.current_item = self.annotation_box.focus()
        self.current_value = self.annotation_box.item(self.current_item, 'values')
        print(self.current_value[0])
        self.annotation_box.item(self.current_item,
                            values=(
                            self.current_value[0], self.dropdown_label.get("1.0", END), self.description_text.get("1.0", END)))

    def selectedValue(self, evt):
        for self.dropvalue in self.dropdown.curselection():
            print(self.dropvalue)
        self.dropdown_label.delete("1.0", "end")
        for self.dropvalue in self.dropdown.curselection():
            self.selected_tag_value = self.dropdown.get(self.dropvalue)
            self.dropdown_label.insert(END, f'{self.selected_tag_value}, ')








